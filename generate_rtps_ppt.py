"""Generate Clean, Academic 6-Slide RTPS Review Presentation in PPTX format.

Clean academic presentation without any competition references.
Focus: Problem, Solution, How it Works, Model & Features, Results, and Graphs with 'Why This Graph'.
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

PROJECT_ROOT = Path(__file__).resolve().parent
ASSETS_DIR = PROJECT_ROOT / "rtps_presentation_assets"
OUTPUT_PPTX = PROJECT_ROOT / "RTPS_Business_Entity_Resolution_Presentation.pptx"

# Modern Academic Palette
NAVY = RGBColor(20, 35, 65)          # Deep Navy
BLUE = RGBColor(30, 110, 200)        # Primary Blue
LIGHT_BG = RGBColor(246, 249, 252)   # Card Fill
WHITE = RGBColor(255, 255, 255)
TEXT_DARK = RGBColor(35, 45, 60)     # Body Text
MUTED = RGBColor(100, 115, 130)      # Secondary
GREEN = RGBColor(20, 140, 80)        # Success / Metric
BORDER = RGBColor(220, 230, 242)     # Subtle card border


def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_slide_header(slide, title, subtitle="REAL-TIME PROJECT SEMINAR (RTPS) | MACHINE LEARNING"):
        box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_sub = tf.paragraphs[0]
        p_sub.text = subtitle.upper()
        p_sub.font.size = Pt(10)
        p_sub.font.bold = True
        p_sub.font.color.rgb = BLUE

        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY

    def add_card(slide, left, top, width, height, bg_color=WHITE, border_color=BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    # =========================================================================
    # SLIDE 1: PROBLEM STATEMENT
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)

    # Hero Banner
    hero = add_card(s1, Inches(0.8), Inches(0.6), Inches(11.733), Inches(2.1), bg_color=NAVY, border_color=None)
    tf_h = hero.text_frame
    tf_h.word_wrap = True
    tf_h.margin_left = Inches(0.5)
    tf_h.margin_top = Inches(0.3)

    p0 = tf_h.paragraphs[0]
    p0.text = "REAL-TIME PROJECT SEMINAR (RTPS) - ML LAB REVIEW"
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(96, 165, 250)

    p1 = tf_h.add_paragraph()
    p1.text = "Multi-Source Business Entity Resolution & Record Linkage"
    p1.font.size = Pt(24)
    p1.font.bold = True
    p1.font.color.rgb = WHITE

    p2 = tf_h.add_paragraph()
    p2.text = "Automated matching of business entities across heterogeneous data sources using Machine Learning"
    p2.font.size = Pt(12)
    p2.font.color.rgb = RGBColor(203, 213, 225)

    # Left Card: What is the Problem?
    c_left = add_card(s1, Inches(0.8), Inches(3.0), Inches(5.7), Inches(3.9))
    tf_l = c_left.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = tf_l.margin_right = Inches(0.35)
    tf_l.margin_top = Inches(0.3)

    p = tf_l.paragraphs[0]
    p.text = "1. What is the Problem?"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY

    points_l = [
        ("Core Task: ", "Identify and link identical business entities that appear across multiple independent databases."),
        ("Canonical vs. Noisy Sources: ", "Source 1 contains clean reference records, while Source 2 and Source 3 contain noisy, unstandardized records."),
        ("Real-World Inconsistencies: ", "Company names and addresses often have spelling typos, missing legal terms, and rearranged words."),
        ("No Unique Primary Keys: ", "Real-world databases do not share common unique IDs (like tax IDs), requiring ML-driven text matching."),
    ]
    for b_title, b_desc in points_l:
        p = tf_l.add_paragraph()
        p.space_before = Pt(8)
        r_b = p.add_run()
        r_b.text = "• " + b_title
        r_b.font.bold = True
        r_b.font.size = Pt(12)
        r_b.font.color.rgb = BLUE
        r_t = p.add_run()
        r_t.text = b_desc
        r_t.font.size = Pt(12)
        r_t.font.color.rgb = TEXT_DARK

    # Right Card: Why is it Difficult?
    c_right = add_card(s1, Inches(6.833), Inches(3.0), Inches(5.7), Inches(3.9))
    tf_r = c_right.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_right = Inches(0.35)
    tf_r.margin_top = Inches(0.3)

    p = tf_r.paragraphs[0]
    p.text = "2. Why is this Challenging?"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY

    points_r = [
        ("Huge Computational Scale: ", "Comparing millions of records directly would require trillions of pairwise comparisons, which is computationally impossible."),
        ("Legal Suffix Variations: ", "Businesses write suffixes differently ('Apple Inc' vs 'Apple Incorporated', 'ABC Corp' vs 'ABC Corporation')."),
        ("Singleton Businesses: ", "Many businesses in Source 1 do not exist in secondary sources and must correctly produce zero matches."),
        ("High Cost of False Matches: ", "Incorrectly merging two different companies causes serious financial and legal errors."),
    ]
    for b_title, b_desc in points_r:
        p = tf_r.add_paragraph()
        p.space_before = Pt(8)
        r_b = p.add_run()
        r_b.text = "• " + b_title
        r_b.font.bold = True
        r_b.font.size = Pt(12)
        r_b.font.color.rgb = GREEN
        r_t = p.add_run()
        r_t.text = b_desc
        r_t.font.size = Pt(12)
        r_t.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 2: PROPOSED SOLUTION
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_header(s2, "Proposed Solution: 2-Stage Hybrid Architecture")

    # Top Explanation Card
    top_c = add_card(s2, Inches(0.8), Inches(1.7), Inches(11.733), Inches(1.0), bg_color=LIGHT_BG)
    tf_t = top_c.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = Inches(0.3)
    tf_t.margin_top = Inches(0.12)
    p = tf_t.paragraphs[0]
    p.text = "Core Solution Strategy"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p_sub = tf_t.add_paragraph()
    p_sub.text = "Instead of comparing every entity with every other entity (O(N²) bottleneck), we use a 2-stage approach: Stage 1 quickly eliminates over 99.8% of obvious non-matches, and Stage 2 runs an ML Ensemble on only the top candidate pairs."
    p_sub.font.size = Pt(11.5)
    p_sub.font.color.rgb = TEXT_DARK

    # Stage 1 Card
    card_s1 = add_card(s2, Inches(0.8), Inches(2.9), Inches(5.7), Inches(4.0))
    tf_s1 = card_s1.text_frame
    tf_s1.word_wrap = True
    tf_s1.margin_left = tf_s1.margin_right = Inches(0.35)
    tf_s1.margin_top = Inches(0.25)

    p = tf_s1.paragraphs[0]
    p.text = "STAGE 1: Candidate Blocker (Fast Filter)"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = BLUE

    s1_bullets = [
        ("Purpose: ", "Quickly find the most likely potential matches without heavy machine learning computations."),
        ("Prefix Indexing: ", "Groups companies sharing the first 4 characters of their name into fast lookup buckets."),
        ("Word Inverted Index: ", "Captures word inversions and common terms across names and street addresses."),
        ("Country Partitioning: ", "Only compares records located in the same country, drastically reducing search space."),
        ("Result: ", "Eliminates >99.8% of irrelevant comparisons and keeps only the top 12 candidate pairs per business."),
    ]
    for b_title, b_desc in s1_bullets:
        p = tf_s1.add_paragraph()
        p.space_before = Pt(7)
        r_b = p.add_run()
        r_b.text = "✔ " + b_title
        r_b.font.bold = True
        r_b.font.size = Pt(11.5)
        r_b.font.color.rgb = NAVY
        r_t = p.add_run()
        r_t.text = b_desc
        r_t.font.size = Pt(11.5)
        r_t.font.color.rgb = TEXT_DARK

    # Stage 2 Card
    card_s2 = add_card(s2, Inches(6.833), Inches(2.9), Inches(5.7), Inches(4.0))
    tf_s2 = card_s2.text_frame
    tf_s2.word_wrap = True
    tf_s2.margin_left = tf_s2.margin_right = Inches(0.35)
    tf_s2.margin_top = Inches(0.25)

    p = tf_s2.paragraphs[0]
    p.text = "STAGE 2: Machine Learning Ensemble Model"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = GREEN

    s2_bullets = [
        ("Purpose: ", "Perform deep, highly accurate comparison on candidate pairs to make the final match decision."),
        ("Feature Extraction: ", "Computes 20 numerical similarity features (spelling similarity, edit distance, token overlap)."),
        ("Model 1 - LightGBM: ", "Fast histogram-based decision tree classifier that handles feature interactions efficiently."),
        ("Model 2 - XGBoost: ", "Gradient boosted tree model with strong regularization to avoid overfitting on rare words."),
        ("Soft Voting: ", "Combines both models (50% LightGBM + 50% XGBoost). Predicts MATCH if probability >= 0.55."),
    ]
    for b_title, b_desc in s2_bullets:
        p = tf_s2.add_paragraph()
        p.space_before = Pt(7)
        r_b = p.add_run()
        r_b.text = "✔ " + b_title
        r_b.font.bold = True
        r_b.font.size = Pt(11.5)
        r_b.font.color.rgb = NAVY
        r_t = p.add_run()
        r_t.text = b_desc
        r_t.font.size = Pt(11.5)
        r_t.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 3: HOW IT WORKS (STEP-BY-STEP WORKFLOW)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_header(s3, "How It Works: Step-by-Step Processing Pipeline")

    steps = [
        ("Step 1: Text Cleaning & Normalization",
         "Convert all text to lowercase, remove punctuation, and standardize legal suffixes (Inc, LLC, Corp, Pvt Ltd) and address abbreviations (St, Rd, Ave).", BLUE),
        ("Step 2: Country Partitioning",
         "Split records by country. This guarantees zero cross-country mismatch errors and allows streaming data in chunks without running out of RAM.", NAVY),
        ("Step 3: Fast Candidate Blocking",
         "Generate top candidate pairs per Source 1 record using prefix lookups and word inverted indexes. Reduces trillions of comparisons to top-12 candidates.", BLUE),
        ("Step 4: Similarity Feature Extraction",
         "Compute 20 pairwise similarity metrics using fast C++ RapidFuzz (Jaro-Winkler, Levenshtein, word overlap, street number matching).", NAVY),
        ("Step 5: ML Ensemble Prediction",
         "Pass features to LightGBM + XGBoost. If probability >= 0.55, classify as TRUE MATCH. If no candidate reaches 0.55, keep as SINGLETON.", GREEN),
        ("Step 6: Output Generation",
         "Produce final formatted output table mapping each Source 1 ID to its verified matching Secondary IDs or empty list if no match exists.", NAVY),
    ]

    for i, (title, desc, color) in enumerate(steps):
        col = i % 3
        row = i // 3
        left = Inches(0.8 + col * 4.0)
        top = Inches(1.8 + row * 2.6)
        width = Inches(3.733)
        height = Inches(2.35)

        card = add_card(s3, left, top, width, height)
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.2)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color

        p2 = tf.add_paragraph()
        p2.space_before = Pt(6)
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 4: FEATURE ENGINEERING & MACHINE LEARNING MODEL
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_header(s4, "Feature Engineering & Machine Learning Model")

    # Left: What Features Did We Create?
    card_f = add_card(s4, Inches(0.8), Inches(1.7), Inches(5.7), Inches(5.2))
    tf_f = card_f.text_frame
    tf_f.word_wrap = True
    tf_f.margin_left = tf_f.margin_right = Inches(0.3)
    tf_f.margin_top = Inches(0.25)

    p = tf_f.paragraphs[0]
    p.text = "Features Extracted per Candidate Pair"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY

    p_sub = tf_f.add_paragraph()
    p_sub.text = "We transform raw business text into 20 numerical similarity metrics:"
    p_sub.font.size = Pt(11)
    p_sub.font.color.rgb = MUTED

    feat_groups = [
        ("1. Business Name Similarities:", [
            "Jaro-Winkler Similarity (handles minor typos with prefix weighting)",
            "Levenshtein Distance (character edit distance)",
            "Sequence Ratio (fuzzy character similarity)",
            "Token Jaccard & Overlap (word-by-word intersection)",
            "Length Difference Ratio (checks if one name is an abbreviation)",
        ]),
        ("2. Address Similarities:", [
            "Street Address Sequence Similarity",
            "Token Jaccard on Address Words",
            "Numeric / Street Number Match (checks if building numbers match)",
            "Postal / Digit Overlap Ratio",
        ]),
        ("3. Source & Compatibility Flags:", [
            "Country Match / Mismatch",
            "Source Indicator Flags (Source 2 vs Source 3)",
        ]),
    ]
    for grp_title, subitems in feat_groups:
        p = tf_f.add_paragraph()
        p.space_before = Pt(6)
        p.text = grp_title
        p.font.bold = True
        p.font.size = Pt(11.5)
        p.font.color.rgb = BLUE
        for item in subitems:
            p_i = tf_f.add_paragraph()
            p_i.text = "  • " + item
            p_i.font.size = Pt(10.5)
            p_i.font.color.rgb = TEXT_DARK

    # Right: The Ensemble Model
    card_m = add_card(s4, Inches(6.833), Inches(1.7), Inches(5.7), Inches(5.2))
    tf_m = card_m.text_frame
    tf_m.word_wrap = True
    tf_m.margin_left = tf_m.margin_right = Inches(0.3)
    tf_m.margin_top = Inches(0.25)

    p = tf_m.paragraphs[0]
    p.text = "ML Ensemble Classifier Architecture"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY

    m_points = [
        ("Why Ensemble Learning? ", "Single models often overfit to specific spelling patterns. Combining two diverse gradient boosted architectures produces smoother and more reliable predictions."),
        ("Model 1: LightGBM ", "Uses histogram-based tree splitting. Extremely fast and handles class imbalance efficiently using automated sample reweighting."),
        ("Model 2: XGBoost ", "Uses depth-wise tree expansion with exact second-order gradients. Provides strong L1 and L2 regularization to prevent false matches."),
        ("Ensemble Equation: ", "Final Probability = 0.5 * P(LightGBM) + 0.5 * P(XGBoost)"),
        ("Decision Cutoff (Threshold = 0.55): ", "Instead of default 0.50, we use 0.55. This slightly conservative threshold ensures high precision and prevents false company mergers."),
    ]
    for b_title, b_desc in m_points:
        p = tf_m.add_paragraph()
        p.space_before = Pt(7)
        r_b = p.add_run()
        r_b.text = "✔ " + b_title
        r_b.font.bold = True
        r_b.font.size = Pt(11)
        r_b.font.color.rgb = GREEN
        r_t = p.add_run()
        r_t.text = b_desc
        r_t.font.size = Pt(11)
        r_t.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 5: OUTPUT & PERFORMANCE RESULTS
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_header(s5, "Model Performance & Output Verification")

    # 4 Metric Cards
    metrics = [
        ("Overall Accuracy", "95.95%", "Percentage of correct match & non-match predictions", GREEN),
        ("Precision", "85.63%", "Out of all predicted matches, 85.6% are genuine matches", BLUE),
        ("Recall (Sensitivity)", "99.60%", "Model successfully captured 3,479 out of 3,493 true matches", NAVY),
        ("ROC-AUC Score", "0.9915", "Demonstrates near-perfect separation between true and false pairs", GREEN),
    ]
    for idx, (m_name, m_val, m_desc, m_color) in enumerate(metrics):
        left = Inches(0.8 + idx * 2.95)
        top = Inches(1.7)
        c_m = add_card(s5, left, top, Inches(2.8), Inches(1.5))
        tf_m = c_m.text_frame
        tf_m.word_wrap = True
        tf_m.margin_left = tf_m.margin_right = Inches(0.2)
        tf_m.margin_top = Inches(0.12)

        p = tf_m.paragraphs[0]
        p.text = m_name
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = MUTED

        p_v = tf_m.add_paragraph()
        p_v.text = m_val
        p_v.font.size = Pt(22)
        p_v.font.bold = True
        p_v.font.color.rgb = m_color

        p_d = tf_m.add_paragraph()
        p_d.text = m_desc
        p_d.font.size = Pt(9)
        p_d.font.color.rgb = TEXT_DARK

    # Bottom Left: Confusion Matrix Numbers
    card_cm = add_card(s5, Inches(0.8), Inches(3.45), Inches(5.7), Inches(3.5))
    tf_cm = card_cm.text_frame
    tf_cm.word_wrap = True
    tf_cm.margin_left = tf_cm.margin_right = Inches(0.3)
    tf_cm.margin_top = Inches(0.2)

    p = tf_cm.paragraphs[0]
    p.text = "Confusion Matrix Audit (14,750 Pairs Tested)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY

    cm_items = [
        ("True Negatives (TN): 10,673 (72.4%)", "Non-matching candidate records correctly identified and rejected."),
        ("True Positives (TP): 3,479 (23.6%)", "Genuine business matches successfully linked across secondary sources."),
        ("False Positives (FP): 584 (4.0%)", "False merges kept low through the 0.55 decision threshold."),
        ("False Negatives (FN): Only 14 (0.1%)", "Only 14 actual matches missed out of 3,493 (99.6% recall)."),
    ]
    for title_cm, desc_cm in cm_items:
        p = tf_cm.add_paragraph()
        p.space_before = Pt(5)
        r1 = p.add_run()
        r1.text = "• " + title_cm + "\n"
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = BLUE
        r2 = p.add_run()
        r2.text = "  " + desc_cm
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_DARK

    # Bottom Right: Real-World Matching Examples
    card_ex = add_card(s5, Inches(6.833), Inches(3.45), Inches(5.7), Inches(3.5))
    tf_ex = card_ex.text_frame
    tf_ex.word_wrap = True
    tf_ex.margin_left = tf_ex.margin_right = Inches(0.3)
    tf_ex.margin_top = Inches(0.2)

    p = tf_ex.paragraphs[0]
    p.text = "Real-World Output Verification Examples"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY

    examples = [
        ("Case 1: Legal Suffix Variation ('Corp' vs 'Corporation')",
         "Source 1: 'Technoblade Solutions Corp' | Austin, TX\nSource 2: 'Technoblade Solutions Corporation' | Austin, TX\nResult: MATCH (Probability: 0.9997) -> CORRECT"),
        ("Case 2: Human Spelling Typo in Business Title",
         "Source 1: 'Omni Retail Logistics' | Chicago, IL\nSource 3: 'Omni Retial Logistcs' | Chicago, IL\nResult: MATCH (Probability: 0.9998) -> CORRECT"),
        ("Case 3: Similar Business Names in Different Locations",
         "Source 1: 'Apex Auto Parts' | Buffalo, NY\nSource 2: 'Apex Auto Body & Repair' | Rochester, NY\nResult: NON-MATCH (Probability: 0.0002) -> CORRECT"),
    ]
    for ex_t, ex_d in examples:
        p = tf_ex.add_paragraph()
        p.space_before = Pt(5)
        r1 = p.add_run()
        r1.text = "✔ " + ex_t + "\n"
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = GREEN
        r2 = p.add_run()
        r2.text = ex_d
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 6: EVALUATION GRAPHS & WHY WE USE THEM
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_header(s6, "Evaluation Graphs & Visual Model Interpretability")

    # Left: Master Dashboard Image
    master_plot = ASSETS_DIR / "rtps_master_dashboard.png"
    if master_plot.exists():
        s6.shapes.add_picture(str(master_plot), Inches(0.8), Inches(1.6), Inches(7.2), Inches(5.4))
    else:
        card_fb = add_card(s6, Inches(0.8), Inches(1.6), Inches(7.2), Inches(5.4))
        tf_fb = card_fb.text_frame
        p = tf_fb.paragraphs[0]
        p.text = "Master Dashboard Image"

    # Right: Why This Graph? (Clear Explanation for Review Panel)
    card_why = add_card(s6, Inches(8.2), Inches(1.6), Inches(4.333), Inches(5.4))
    tf_w = card_why.text_frame
    tf_w.word_wrap = True
    tf_w.margin_left = tf_w.margin_right = Inches(0.25)
    tf_w.margin_top = Inches(0.2)

    p = tf_w.paragraphs[0]
    p.text = "Why These Graphs?"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY

    why_points = [
        ("1. Confusion Matrix (Top-Left):",
         "Why it matters: Proves that the model works on both classes. It confirms 10,673 non-matches were rejected and only 14 true matches were missed (99.6% recall)."),
        ("2. ROC Curve (Top-Right):",
         "Why it matters: Shows the trade-off between True Positives and False Positives. An AUC of 0.9915 proves that the model separates matches from non-matches with near-perfect reliability."),
        ("3. Feature Importance (Bottom-Left):",
         "Why it matters: Provides interpretability. It shows that name length difference (33.9%) and string similarity (18.9%) are the most critical factors for matching."),
        ("4. Threshold Optimization (Bottom-Right):",
         "Why it matters: Explains why 0.55 cutoff was selected. It balances precision and recall to strictly protect against false company mergers."),
    ]
    for w_title, w_desc in why_points:
        p = tf_w.add_paragraph()
        p.space_before = Pt(7)
        r1 = p.add_run()
        r1.text = w_title + "\n"
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = BLUE
        r2 = p.add_run()
        r2.text = w_desc
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_DARK

    save_success = False
    target_path = OUTPUT_PPTX
    try:
        prs.save(str(target_path))
        save_success = True
    except PermissionError:
        alt_path = PROJECT_ROOT / "Business_Entity_Resolution_RTPS_Review.pptx"
        try:
            prs.save(str(alt_path))
            target_path = alt_path
            save_success = True
        except Exception as e:
            print(f"[ERROR] Could not save presentation: {e}")

    if save_success:
        print(f"[SUCCESS] Academic RTPS Presentation generated: {target_path}")
        print(f"Total Slides: {len(prs.slides)}")


if __name__ == "__main__":
    create_presentation()
